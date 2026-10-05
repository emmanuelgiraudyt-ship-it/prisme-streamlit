import io
from email import message_from_bytes, policy

import pytest
import requests
from docx import Document

from prisme import exports

CONTENT = "# Objet\n\nUn **paragraphe** important.\n- point un\n- point deux\n1. étape\n\nConclusion."


def test_docx_structure():
    data = exports.to_docx_bytes("Note test", CONTENT, author="Cabinet")
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    assert "Note test" in texts and "Objet" in texts
    assert any(p.text == "paragraphe" or "paragraphe" in p.text for p in doc.paragraphs)
    assert any(p.style.name == "List Bullet" for p in doc.paragraphs)
    assert exports.MENTION_IA in texts
    assert doc.core_properties.author == "Cabinet"


def test_eml_is_unsent_draft():
    msg = message_from_bytes(exports.to_eml_bytes("Objet accentué é", "Corps", to="a@b.fr"), policy=policy.default)
    assert msg["X-Unsent"] == "1" and msg["Subject"] == "Objet accentué é" and msg["To"] == "a@b.fr"
    assert "Corps" in msg.get_content()


def test_extract_notion_id():
    raw = "0123456789abcdef0123456789abcdef"
    assert exports.extract_notion_id(f"https://www.notion.so/Ma-page-{raw}") == raw
    assert exports.extract_notion_id("01234567-89ab-cdef-0123-456789abcdef") == raw
    assert exports.extract_notion_id("rien") is None


def test_notion_blocks():
    blocks = exports.to_notion_blocks("# Titre\nTexte\n- puce\n" + "x" * 4000)
    kinds = [b["type"] for b in blocks]
    assert kinds[:3] == ["heading_1", "paragraph", "bulleted_list_item"]
    assert kinds.count("paragraph") == 1 + 3  # 4000 caractères découpés en 3 blocs
    assert all(len(b[b["type"]]["rich_text"][0]["text"]["content"]) <= 1900 for b in blocks)


class Resp:
    def __init__(self, data, status=200):
        self._data, self.status_code, self.text = data, status, str(data)

    def json(self):
        return self._data


def test_notion_page_creation_batches(monkeypatch):
    monkeypatch.setenv("NOTION_TOKEN", "tok")
    monkeypatch.setenv("NOTION_PARENT_PAGE_ID", "0123456789abcdef0123456789abcdef")
    posts, patches = [], []
    monkeypatch.setattr(
        requests,
        "post",
        lambda url, headers, json, timeout: posts.append(json) or Resp({"id": "p1", "url": "https://notion.so/p1"}),
    )
    monkeypatch.setattr(requests, "patch", lambda url, headers, json, timeout: patches.append(json) or Resp({}))
    url = exports.notion_create_page("Titre", "\n".join(f"ligne {i}" for i in range(230)))
    assert url == "https://notion.so/p1"
    assert posts[0]["parent"] == {"page_id": "0123456789abcdef0123456789abcdef"}
    assert len(posts[0]["children"]) == 100 and [len(p["children"]) for p in patches] == [100, 30]


def test_notion_database_uses_title_property(monkeypatch):
    monkeypatch.setenv("NOTION_TOKEN", "tok")
    monkeypatch.setenv("NOTION_DATABASE_ID", "0123456789abcdef0123456789abcdef")
    sent = []
    monkeypatch.setattr(
        requests,
        "get",
        lambda url, headers, timeout: Resp({"properties": {"Séance": {"type": "title"}, "X": {"type": "text"}}}),
    )
    monkeypatch.setattr(
        requests, "post", lambda url, headers, json, timeout: sent.append(json) or Resp({"id": "p", "url": "u"})
    )
    exports.notion_create_page("T", "texte")
    assert "Séance" in sent[0]["properties"] and "database_id" in sent[0]["parent"]


def test_notion_error_is_explicit(monkeypatch):
    monkeypatch.setenv("NOTION_TOKEN", "tok")
    monkeypatch.setenv("NOTION_PARENT_PAGE_ID", "0123456789abcdef0123456789abcdef")
    monkeypatch.setattr(requests, "post", lambda *a, **k: Resp({"message": "page introuvable"}, 404))
    with pytest.raises(exports.ExportError, match="page introuvable"):
        exports.notion_create_page("T", "x")


def test_notion_not_configured():
    assert not exports.notion_configured()
    with pytest.raises(exports.ExportError):
        exports.notion_create_page("T", "x")
